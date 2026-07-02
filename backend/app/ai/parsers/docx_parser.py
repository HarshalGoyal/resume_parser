from typing import Dict

from docx import Document

from app.models.parsed_document import ParsedDocument
from app.models.document_line import DocumentLine
from app.models.document_section import DocumentSection
from app.core.logging import AppLogger
from app.utils.text_utils import (
    collapse_whitespace,
    split_layout_segments,
    looks_like_section_header,
)


class DOCXParser:
    """Structural parser for .docx resumes.

    Produces the same ``ParsedDocument`` shape as ``PDFParser`` so the
    downstream enrichment service is agnostic to the source format. Header
    detection uses generic, template-independent signals (Word heading styles,
    a broad section-keyword vocabulary, and the all-caps convention) rather than
    font/bold cues, which are unreliable across templates.
    """

    def __init__(self):
        self.logger = AppLogger("DOCXParser")
        self.logger.info("Initialized DOCXParser for handling DOCX parsing logic.")

    async def parse(self, file_path: str) -> ParsedDocument:
        self.logger.info(f"Parsing DOCX at path: {file_path}")
        document = Document(file_path)

        parsed_lines: list[DocumentLine] = []
        for paragraph in document.paragraphs:
            raw = paragraph.text
            text = collapse_whitespace(raw)
            if not text:
                continue
            parsed_lines.append(
                DocumentLine(
                    text=text,
                    segments=split_layout_segments(raw),
                    bold=self.__is_bold(paragraph),
                    is_header=self.__is_header(paragraph, text),
                )
            )

        sections = self.__build_sections(parsed_lines)
        self.logger.info(f"parsed {len(sections)} sections from DOCX")
        return ParsedDocument(metadata={"paragraphs": len(parsed_lines)}, sections=sections)

    def __is_header(self, paragraph, text: str) -> bool:
        style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style_name.startswith("heading") or style_name == "title":
            return True
        # Generic conventions (keyword / all-caps), never font/bold-only — a bold
        # project or job title must not be mistaken for a section header.
        return looks_like_section_header(text)

    def __is_bold(self, paragraph) -> bool:
        runs = [r for r in paragraph.runs if r.text.strip()]
        if not runs:
            return False
        return any(bool(r.bold) for r in runs)

    def __build_sections(self, parsed_lines) -> Dict[str, DocumentSection]:
        sections: dict[str, list[DocumentLine]] = {"HEADER": []}
        current_section = "HEADER"

        for line in parsed_lines:
            if line.is_header:
                current_section = line.text
                sections.setdefault(current_section, [])
                continue
            sections[current_section].append(line)

        return {
            name: DocumentSection(
                name=name,
                content="\n".join(line.text for line in lines),
                lines=lines,
            )
            for name, lines in sections.items()
        }
